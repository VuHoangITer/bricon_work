"""Thanh toán công nợ — NVKD khai báo mỗi lần khách chuyển tiền (tên khách,
mã đơn, ảnh lệnh chuyển tiền, ghi chú); Quản lý bộ phận / Ban giám đốc /
Quản trị tra cứu lại.

Quyền:
- Khai báo: Ban giám đốc + Quản trị, và nhân viên có CHỨC VỤ được tick ở ô "Chức vụ được khai báo"
  (Quản trị chỉnh ngay trên trang này). Chưa tick gì thì mặc định mọi chức
  vụ có chữ "kinh doanh" / "NVKD" / "sale" trong tên.
- Xem: người khai báo xem phiếu của mình; Quản lý bộ phận xem phiếu của
  người trong bộ phận; Ban giám đốc + Quản trị xem tất cả.
- Sửa: người khai báo sửa phiếu của mình trong 24 giờ; Quản trị sửa mọi lúc.
- Xoá: chỉ Quản trị.
"""
import unicodedata
from datetime import date, datetime, timedelta
from io import BytesIO

from flask import (Blueprint, abort, flash, jsonify, redirect, render_template, request, send_file,
                   url_for)
from flask_login import current_user, login_required
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

import services
from extensions import db
from models import (AnhCongNo, ChucVu, LoaiDinhKem, NguoiDung, ThanhToanCongNo, VaiTro,
                    gio_vn_hien_tai, ngay_vn_hien_tai)

bp = Blueprint("cong_no", __name__, url_prefix="/cong-no")

KHOA_CHUC_VU = "cong_no_chuc_vu"
GIO_DUOC_SUA = 24


# ---------------------------------------------------------------- tiện ích
def khong_dau(s: str | None) -> str:
    """Bỏ dấu tiếng Việt, chữ thường, gộp khoảng trắng — dùng cho tìm kiếm
    và gộp ghi chú."""
    s = unicodedata.normalize("NFKD", s or "")
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = s.replace("đ", "d").replace("Đ", "D").lower()
    return " ".join(s.split())


def _gon(s: str | None) -> str:
    return " ".join((s or "").split())


def _chuc_vu_khai_bao_ids() -> set[int]:
    gia_tri = services.lay_cai_dat(KHOA_CHUC_VU) or ""
    ids = {int(x) for x in gia_tri.split(",") if x.strip().isdigit()}
    if ids:
        return ids
    tu_khoa = ("kinh doanh", "nvkd", "sale")
    return {cv.id for cv in ChucVu.query.all() if any(t in khong_dau(cv.ten) for t in tu_khoa)}


def duoc_khai_bao(nd) -> bool:
    """Ban giám đốc + Quản trị luôn được khai báo; nhân viên khác theo chức vụ."""
    if not (nd and nd.is_authenticated):
        return False
    if nd.vai_tro in (VaiTro.ADMIN, VaiTro.SEP):
        return True
    return bool(nd.chuc_vu_id and nd.chuc_vu_id in _chuc_vu_khai_bao_ids())


def duoc_vao(nd) -> bool:
    """Ai thấy mục Thanh toán công nợ trên menu."""
    if not (nd and nd.is_authenticated):
        return False
    if nd.vai_tro in (VaiTro.ADMIN, VaiTro.SEP) or duoc_khai_bao(nd):
        return True
    if nd.vai_tro == VaiTro.QUAN_LY and nd.ds_bo_phan_id:
        ids = _chuc_vu_khai_bao_ids()
        return bool(ids) and db.session.query(NguoiDung.id).filter(
            NguoiDung.thuoc_bo_phan(nd.ds_bo_phan_id), NguoiDung.chuc_vu_id.in_(ids)).first() is not None
    return False


def duoc_sua(nd, cn: ThanhToanCongNo) -> bool:
    if nd.la_admin_thuan:
        return True
    return (cn.nguoi_tao_id == nd.id and cn.tao_luc is not None
            and gio_vn_hien_tai() - cn.tao_luc <= timedelta(hours=GIO_DUOC_SUA))


def duoc_xem(nd, cn: ThanhToanCongNo) -> bool:
    if cn.nguoi_tao_id == nd.id:
        return True
    return nd.la_quan_ly and nd.trong_pham_vi(cn.nguoi_tao)


def _q_pham_vi():
    """Query phiếu trong phạm vi người đang đăng nhập được xem."""
    q = ThanhToanCongNo.query
    u = current_user
    if u.vai_tro in (VaiTro.ADMIN, VaiTro.SEP):
        return q
    if u.vai_tro == VaiTro.QUAN_LY and u.ds_bo_phan_id:
        return (q.join(NguoiDung, ThanhToanCongNo.nguoi_tao_id == NguoiDung.id)
                .filter(db.or_(NguoiDung.thuoc_bo_phan(u.ds_bo_phan_id),
                               ThanhToanCongNo.nguoi_tao_id == u.id)))
    return q.filter(ThanhToanCongNo.nguoi_tao_id == u.id)


def _bat_buoc_vao():
    if not duoc_vao(current_user):
        abort(403)


def _doc_ngay(s: str | None) -> date | None:
    try:
        return datetime.strptime(s, "%Y-%m-%d").date() if s else None
    except ValueError:
        return None


def _bo_loc():
    """Đọc bộ lọc từ URL — dùng chung cho danh sách và xuất Excel."""
    f = {
        "q": _gon(request.args.get("q")),
        "tu": _doc_ngay(request.args.get("tu")),
        "den": _doc_ngay(request.args.get("den")),
        "gc": khong_dau(request.args.get("gc")),
        "nguoi": request.args.get("nguoi", type=int),
    }
    q = _q_pham_vi()
    if f["q"]:
        q = q.filter(ThanhToanCongNo.tim_kiem.like(f"%{khong_dau(f['q'])}%"))
    if f["tu"]:
        q = q.filter(ThanhToanCongNo.tao_luc >= datetime.combine(f["tu"], datetime.min.time()))
    if f["den"]:
        q = q.filter(ThanhToanCongNo.tao_luc < datetime.combine(f["den"] + timedelta(days=1),
                                                                datetime.min.time()))
    if f["gc"]:
        q = q.filter(ThanhToanCongNo.ghi_chu_chuan == f["gc"])
    if f["nguoi"] and current_user.la_quan_ly:
        q = q.filter(ThanhToanCongNo.nguoi_tao_id == f["nguoi"])
    return f, q.order_by(ThanhToanCongNo.tao_luc.desc(), ThanhToanCongNo.id.desc())


def _ds_ghi_chu():
    """Các ghi chú đã dùng trong phạm vi xem — [(chuẩn, nhãn, số phiếu)],
    nhiều nhất trước. Ghi chú chỉ khác hoa/thường/dấu cách gộp làm một."""
    sub = _q_pham_vi().with_entities(ThanhToanCongNo.id).subquery()
    rows = (db.session.query(ThanhToanCongNo.ghi_chu_chuan, db.func.max(ThanhToanCongNo.ghi_chu),
                             db.func.count(ThanhToanCongNo.id))
            .filter(ThanhToanCongNo.id.in_(db.select(sub.c.id)))
            .group_by(ThanhToanCongNo.ghi_chu_chuan)
            .order_by(db.func.count(ThanhToanCongNo.id).desc()).limit(200).all())
    return [(c, nhan, n) for c, nhan, n in rows]


def _ds_ten_kh(gioi_han: int = 500) -> list[str]:
    sub = _q_pham_vi().with_entities(ThanhToanCongNo.id).subquery()
    rows = (db.session.query(ThanhToanCongNo.ten_kh, db.func.max(ThanhToanCongNo.tao_luc))
            .filter(ThanhToanCongNo.id.in_(db.select(sub.c.id)))
            .group_by(ThanhToanCongNo.ten_kh)
            .order_by(db.func.max(ThanhToanCongNo.tao_luc).desc()).limit(gioi_han).all())
    return [r[0] for r in rows]


def _anh_hop_le(cac_anh) -> bool:
    return all(services.phan_loai(f.filename, f.mimetype) == LoaiDinhKem.ANH for f in cac_anh)


def _gan_truong(cn: ThanhToanCongNo, ten_kh: str, ma_don: str, ghi_chu: str):
    cn.ten_kh = ten_kh[:200]
    cn.ma_don = ma_don[:50]
    cn.ghi_chu = ghi_chu[:255]
    cn.ghi_chu_chuan = khong_dau(ghi_chu)[:255]
    cn.tim_kiem = khong_dau(f"{ten_kh} {ma_don}")[:500]


def _doc_form():
    return (_gon(request.form.get("ten_kh")).upper(), _gon(request.form.get("ma_don")),
            _gon(request.form.get("ghi_chu")))


# ------------------------------------------------------------------ routes
@bp.route("/")
@login_required
def danh_sach():
    _bat_buoc_vao()
    f, q = _bo_loc()
    phan_trang = q.paginate(page=request.args.get("trang", 1, type=int), per_page=30, error_out=False)
    nhan_vien = []
    if current_user.la_quan_ly:
        ids = {r[0] for r in _q_pham_vi().with_entities(ThanhToanCongNo.nguoi_tao_id).distinct()}
        nhan_vien = (NguoiDung.query.filter(NguoiDung.id.in_(ids)).order_by(NguoiDung.ho_ten).all()
                     if ids else [])
    hom_nay = ngay_vn_hien_tai()
    nhanh = {
        "hom_nay": (hom_nay, hom_nay),
        "tuan_nay": (hom_nay - timedelta(days=hom_nay.weekday()), hom_nay),
        "thang_nay": (hom_nay.replace(day=1), hom_nay),
    }
    tham_so = {k: v for k, v in {
        "q": f["q"], "tu": f["tu"].isoformat() if f["tu"] else "",
        "den": f["den"].isoformat() if f["den"] else "",
        "gc": request.args.get("gc") or "", "nguoi": f["nguoi"] or ""}.items() if v}
    ds_chuc_vu = ChucVu.query.order_by(ChucVu.ten).all() if current_user.la_admin_thuan else []
    zalo_nhom = services.lay_cai_dat("cong_no_group_id") if current_user.la_admin_thuan else None
    return render_template(
        "cong_no_list.html", phan_trang=phan_trang, ban_ghi=phan_trang.items, f=f,
        tham_so=tham_so, nhan_vien=nhan_vien, ds_ghi_chu=_ds_ghi_chu(), nhanh=nhanh,
        duoc_khai_bao=duoc_khai_bao(current_user), duoc_sua=duoc_sua,
        ds_chuc_vu=ds_chuc_vu, chuc_vu_chon=_chuc_vu_khai_bao_ids(),
        chuc_vu_da_luu=bool(services.lay_cai_dat(KHOA_CHUC_VU)), zalo_nhom=zalo_nhom,
    )


@bp.route("/goi-y")
@login_required
def goi_y():
    """Gợi ý khi gõ ô tìm kiếm: tên khách + mã đơn khớp (không dấu)."""
    if not duoc_vao(current_user):
        return jsonify([])
    tu = khong_dau(request.args.get("q"))
    if not tu:
        return jsonify([])
    q = _q_pham_vi().filter(ThanhToanCongNo.tim_kiem.like(f"%{tu}%"))
    ket_qua, da_co = [], set()
    for cn in q.order_by(ThanhToanCongNo.tao_luc.desc()).limit(60).all():
        if tu in khong_dau(cn.ten_kh) and ("kh", cn.ten_kh) not in da_co:
            da_co.add(("kh", cn.ten_kh))
            ket_qua.append({"loai": "kh", "text": cn.ten_kh})
        if tu in khong_dau(cn.ma_don) and ("ma", cn.ma_don) not in da_co:
            da_co.add(("ma", cn.ma_don))
            ket_qua.append({"loai": "ma", "text": cn.ma_don, "phu": cn.ten_kh})
        if len(ket_qua) >= 10:
            break
    return jsonify(ket_qua)


@bp.route("/moi")
@login_required
def moi():
    if not duoc_khai_bao(current_user):
        abort(403)
    return render_template("cong_no_form.html", cn=None, ds_ten_kh=_ds_ten_kh(),
                           ds_ghi_chu=[n for _, n, _ in _ds_ghi_chu()])


@bp.route("/luu", methods=["POST"])
@login_required
def luu():
    if not duoc_khai_bao(current_user):
        abort(403)
    ten_kh, ma_don, ghi_chu = _doc_form()
    cac_anh = [f for f in request.files.getlist("anh") if f and f.filename]

    thieu = [ten for ten, gt in (("Tên khách hàng", ten_kh), ("Mã đơn hàng", ma_don),
                                 ("Hình ảnh thanh toán", cac_anh), ("Ghi chú", ghi_chu)) if not gt]
    if thieu:
        flash("Chưa nhập: " + ", ".join(thieu) + ".", "error")
        return redirect(url_for("cong_no.moi"))
    if not _anh_hop_le(cac_anh):
        flash("Có file không phải ảnh — chỉ gửi ảnh lệnh chuyển tiền.", "error")
        return redirect(url_for("cong_no.moi"))

    # Chặn bấm Lưu 2 lần: cùng người, cùng khách, cùng mã trong 3 phút
    trung_vua_xong = ThanhToanCongNo.query.filter(
        ThanhToanCongNo.nguoi_tao_id == current_user.id, ThanhToanCongNo.ma_don == ma_don[:50],
        ThanhToanCongNo.ten_kh == ten_kh[:200],
        ThanhToanCongNo.tao_luc >= gio_vn_hien_tai() - timedelta(minutes=3)).first()
    if trung_vua_xong:
        flash(f"Phiếu {ma_don} – {ten_kh} vừa được lưu rồi, không lưu thêm lần nữa.", "error")
        return redirect(url_for("cong_no.chi_tiet", id=trung_vua_xong.id))

    da_co_ma = ThanhToanCongNo.query.filter(ThanhToanCongNo.ma_don == ma_don[:50]).count()

    cn = ThanhToanCongNo(nguoi_tao_id=current_user.id)
    _gan_truong(cn, ten_kh, ma_don, ghi_chu)
    db.session.add(cn)
    db.session.flush()
    for f in cac_anh:
        duong_dan, _ = services.luu_file(f, "cong-no")
        db.session.add(AnhCongNo(cong_no_id=cn.id, duong_dan=duong_dan))
    db.session.commit()

    services.bao_cong_no_moi(cn)
    db.session.commit()  # commit LogZalo vừa ghi

    flash(f"Đã lưu: {ten_kh} – {ma_don} ({len(cac_anh)} ảnh).", "success")
    if da_co_ma:
        flash(f"Lưu ý: mã {ma_don} đã có {da_co_ma} phiếu thanh toán trước đó.", "warning")
    return redirect(url_for("cong_no.danh_sach"))


@bp.route("/<int:id>")
@login_required
def chi_tiet(id):
    _bat_buoc_vao()
    cn = db.session.get(ThanhToanCongNo, id) or abort(404)
    if not duoc_xem(current_user, cn):
        abort(403)
    cung_ma = (_q_pham_vi().filter(ThanhToanCongNo.ma_don == cn.ma_don, ThanhToanCongNo.id != cn.id)
               .order_by(ThanhToanCongNo.tao_luc.desc()).all())
    return render_template("cong_no_detail.html", cn=cn, duoc_sua=duoc_sua(current_user, cn),
                           cung_ma=cung_ma, gio_duoc_sua=GIO_DUOC_SUA)


@bp.route("/<int:id>/sua", methods=["GET", "POST"])
@login_required
def sua(id):
    cn = db.session.get(ThanhToanCongNo, id) or abort(404)
    if not duoc_sua(current_user, cn):
        flash(f"Chỉ sửa được phiếu của mình trong {GIO_DUOC_SUA} giờ sau khi khai báo — "
              "quá hạn thì nhờ Quản trị sửa.", "error")
        return redirect(url_for("cong_no.chi_tiet", id=id) if duoc_xem(current_user, cn)
                        else url_for("cong_no.danh_sach"))

    if request.method == "GET":
        return render_template("cong_no_form.html", cn=cn, ds_ten_kh=_ds_ten_kh(),
                               ds_ghi_chu=[n for _, n, _ in _ds_ghi_chu()])

    ten_kh, ma_don, ghi_chu = _doc_form()
    anh_moi = [f for f in request.files.getlist("anh") if f and f.filename]
    xoa_ids = {int(x) for x in request.form.getlist("xoa_anh") if x.isdigit()}
    anh_giu = [a for a in cn.anh if a.id not in xoa_ids]

    thieu = [ten for ten, gt in (("Tên khách hàng", ten_kh), ("Mã đơn hàng", ma_don),
                                 ("Hình ảnh thanh toán", anh_giu or anh_moi), ("Ghi chú", ghi_chu))
             if not gt]
    if thieu:
        flash("Chưa nhập: " + ", ".join(thieu) + ".", "error")
        return redirect(url_for("cong_no.sua", id=id))
    if not _anh_hop_le(anh_moi):
        flash("Có file không phải ảnh — chỉ gửi ảnh lệnh chuyển tiền.", "error")
        return redirect(url_for("cong_no.sua", id=id))

    thay_doi = []
    for nhan, cu, moi_ in (("Tên KH", cn.ten_kh, ten_kh[:200]), ("Mã", cn.ma_don, ma_don[:50]),
                           ("Ghi chú", cn.ghi_chu, ghi_chu[:255])):
        if cu != moi_:
            thay_doi.append(f"{nhan}: {cu} → {moi_}")
    so_xoa = len(cn.anh) - len(anh_giu)
    if so_xoa:
        thay_doi.append(f"bỏ {so_xoa} ảnh")
    if anh_moi:
        thay_doi.append(f"thêm {len(anh_moi)} ảnh")
    if not thay_doi:
        flash("Không có gì thay đổi.", "info")
        return redirect(url_for("cong_no.chi_tiet", id=id))

    _gan_truong(cn, ten_kh, ma_don, ghi_chu)
    for a in list(cn.anh):
        if a.id in xoa_ids:
            services._xoa_file_tuong_doi(a.duong_dan)
            cn.anh.remove(a)
    for f in anh_moi:
        duong_dan, _ = services.luu_file(f, "cong-no")
        cn.anh.append(AnhCongNo(duong_dan=duong_dan))
    bay_gio = gio_vn_hien_tai()
    cn.sua_luc = bay_gio
    dong = f"{bay_gio.strftime('%H:%M %d/%m/%Y')} — {current_user.ho_ten}: " + "; ".join(thay_doi)
    cn.lich_su = (cn.lich_su + "\n" if cn.lich_su else "") + dong
    db.session.commit()
    flash("Đã lưu thay đổi.", "success")
    return redirect(url_for("cong_no.chi_tiet", id=id))


@bp.route("/<int:id>/xoa", methods=["POST"])
@login_required
def xoa(id):
    if not current_user.la_admin_thuan:
        abort(403)
    cn = db.session.get(ThanhToanCongNo, id) or abort(404)
    for a in cn.anh:
        services._xoa_file_tuong_doi(a.duong_dan)
    db.session.delete(cn)
    db.session.commit()
    flash(f"Đã xoá phiếu {cn.ma_don} – {cn.ten_kh}.", "success")
    return redirect(url_for("cong_no.danh_sach"))


@bp.route("/cai-dat", methods=["POST"])
@login_required
def cai_dat():
    """Quản trị chọn chức vụ nào được khai báo (NVKD, NVKD Online...)."""
    if not current_user.la_admin_thuan:
        abort(403)
    ids = sorted({int(x) for x in request.form.getlist("chuc_vu") if x.isdigit()
                  and db.session.get(ChucVu, int(x))})
    services.dat_cai_dat(KHOA_CHUC_VU, ",".join(map(str, ids)))
    db.session.commit()
    flash("Đã lưu chức vụ được khai báo." if ids else
          "Đã bỏ chọn — quay về mặc định (chức vụ có chữ \"kinh doanh\" / NVKD / sale).", "success")
    return redirect(url_for("cong_no.danh_sach"))


@bp.route("/xuat-excel")
@login_required
def xuat_excel():
    if not current_user.la_quan_ly:
        abort(403)
    _, q = _bo_loc()
    wb = Workbook()
    ws = wb.active
    ws.title = "Thanh toán công nợ"
    tieu_de = ["STT", "Ngày giờ", "Tên khách hàng", "Mã đơn hàng", "Ghi chú", "Người khai báo", "Số ảnh"]
    ws.append(tieu_de)
    for o in ws[1]:
        o.font = Font(bold=True, color="FFFFFF")
        o.fill = PatternFill("solid", fgColor="1F4E79")
        o.alignment = Alignment(horizontal="center", vertical="center")
    for i, cn in enumerate(q.all(), 1):
        ws.append([i, cn.tao_luc.strftime("%d/%m/%Y %H:%M") if cn.tao_luc else "", cn.ten_kh,
                   cn.ma_don, cn.ghi_chu, cn.nguoi_tao.ho_ten if cn.nguoi_tao else "", len(cn.anh)])
    for cot, rong in zip("ABCDEFG", (6, 18, 32, 16, 34, 24, 8)):
        ws.column_dimensions[cot].width = rong
    ws.freeze_panes = "A2"
    tep = BytesIO()
    wb.save(tep)
    tep.seek(0)
    return send_file(tep, as_attachment=True,
                     download_name=f"thanh-toan-cong-no-{ngay_vn_hien_tai().strftime('%Y-%m-%d')}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")