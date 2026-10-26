"""Biên bản phạt / vi phạm.

Quy trình: Quản lý bộ phận / Ban giám đốc / Quản trị LẬP biên bản (chọn lỗi
từ danh mục -> tự điền mức phạt, đính kèm bằng chứng, ký tên) -> báo Zalo vào
nhóm của nhân viên -> nhân viên vào web KÝ XÁC NHẬN (ghi ý kiến giải trình
nếu có). Nhân viên không chịu ký thì người lập bấm "Ghi nhận NV không ký" —
biên bản vẫn có hiệu lực. Mỗi bước sinh lại PDF.

Quyền:
- Lập: Quản lý bộ phận (cho nhân viên trong bộ phận mình), Ban giám đốc +
  Quản trị (cho mọi người). Không tự lập cho chính mình.
- Xem: nhân viên bị lập, người lập, quản lý trong phạm vi, Ban giám đốc, Quản trị.
- Huỷ: Ban giám đốc / Quản trị bất cứ lúc nào; người lập khi NV chưa ký.
- Xoá vĩnh viễn: chỉ Quản trị.
- Danh mục lỗi: Ban giám đốc + Quản trị.
"""
import base64
import binascii
from datetime import date, datetime, timedelta
from io import BytesIO

from flask import Blueprint, abort, flash, redirect, render_template, request, send_file, url_for
from flask_login import current_user, login_required
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

import services
from extensions import db
from models import (BienBanPhat, DinhKemBienBan, LoaiDinhKem, LoiViPham, NguoiDung, TrangThaiBienBan,
                    VaiTro, gio_vn_hien_tai, ngay_vn_hien_tai)

bp = Blueprint("bien_ban", __name__, url_prefix="/bien-ban")


# ---------------------------------------------------------------- quyền
def duoc_lap_cho(nv: NguoiDung | None) -> bool:
    u = current_user
    if nv is None or not u.la_quan_ly or nv.id == u.id or not nv.dang_hoat_dong:
        return False
    if not u.trong_pham_vi(nv):
        return False
    # Quản lý bộ phận chỉ lập cho NHÂN VIÊN (không lập cho quản lý ngang cấp)
    return u.la_admin_sep or nv.vai_tro == VaiTro.NHAN_VIEN


def duoc_xem(bb: BienBanPhat) -> bool:
    u = current_user
    if u.id in (bb.nguoi_vi_pham_id, bb.nguoi_lap_id):
        return True
    return u.la_quan_ly and u.trong_pham_vi(bb.nguoi_vi_pham)


def duoc_huy(bb: BienBanPhat) -> bool:
    if bb.trang_thai == TrangThaiBienBan.DA_HUY:
        return False
    if current_user.la_admin_sep:
        return True
    return bb.nguoi_lap_id == current_user.id and bb.trang_thai == TrangThaiBienBan.CHO_KY


def duoc_ghi_khong_ky(bb: BienBanPhat) -> bool:
    return bb.trang_thai == TrangThaiBienBan.CHO_KY and (
        bb.nguoi_lap_id == current_user.id or current_user.la_admin_sep)


def _q_pham_vi():
    u = current_user
    q = BienBanPhat.query
    if u.la_admin_sep:
        return q
    if u.vai_tro == VaiTro.QUAN_LY and u.ds_bo_phan_id:
        return (q.join(NguoiDung, BienBanPhat.nguoi_vi_pham_id == NguoiDung.id)
                .filter(db.or_(NguoiDung.thuoc_bo_phan(u.ds_bo_phan_id),
                               BienBanPhat.nguoi_lap_id == u.id,
                               BienBanPhat.nguoi_vi_pham_id == u.id)))
    return q.filter(BienBanPhat.nguoi_vi_pham_id == u.id)


def _nhan_vien_lap_duoc() -> list[NguoiDung]:
    ds = NguoiDung.query.filter_by(dang_hoat_dong=True).order_by(NguoiDung.ho_ten).all()
    return [n for n in ds if duoc_lap_cho(n)]


# ---------------------------------------------------------------- tiện ích
def _doc_chu_ky(ten_truong: str = "chu_ky") -> bytes | None:
    raw = request.form.get(ten_truong) or ""
    if "," not in raw:
        return None
    try:
        du_lieu = base64.b64decode(raw.split(",", 1)[1])
    except (ValueError, binascii.Error):
        return None
    return du_lieu if len(du_lieu) >= 200 else None


def _luu_chu_ky(png: bytes) -> str:
    duong_dan, _ = services.luu_bytes(services.cat_anh_chu_ky(png), "chu-ky-bien-ban.png", "bien-ban")
    return duong_dan


def _doc_tien(s: str | None) -> int | None:
    s = (s or "").strip().replace(".", "").replace(",", "").replace(" ", "").replace("đ", "")
    if not s:
        return 0
    return int(s) if s.isdigit() else None


def _sinh_pdf(bb: BienBanPhat):
    cu = bb.duong_dan_pdf
    bb.duong_dan_pdf = services.luu_pdf_bien_ban(services.tao_pdf_bien_ban(bb))
    if cu:
        services._xoa_file_tuong_doi(cu)


def _bo_loc():
    f = {
        "thang": (request.args.get("thang") or "").strip(),
        "nguoi": request.args.get("nguoi", type=int),
        "tt": request.args.get("tt") or "",
    }
    q = _q_pham_vi()
    if f["thang"]:
        try:
            dau = datetime.strptime(f["thang"], "%Y-%m").date()
            cuoi = date(dau.year + (dau.month == 12), dau.month % 12 + 1, 1)
            q = q.filter(BienBanPhat.ngay_vi_pham >= dau, BienBanPhat.ngay_vi_pham < cuoi)
        except ValueError:
            f["thang"] = ""
    if f["nguoi"] and current_user.la_quan_ly:
        q = q.filter(BienBanPhat.nguoi_vi_pham_id == f["nguoi"])
    if f["tt"] in TrangThaiBienBan.NHAN:
        q = q.filter(BienBanPhat.trang_thai == f["tt"])
    return f, q.order_by(BienBanPhat.ngay_vi_pham.desc(), BienBanPhat.id.desc())


# ------------------------------------------------------------------ routes
@bp.route("/")
@login_required
def danh_sach():
    f, q = _bo_loc()
    ds = q.all()
    tong_theo_nv: dict[int, dict] = {}
    for bb in ds:
        o = tong_theo_nv.setdefault(bb.nguoi_vi_pham_id, {"nv": bb.nguoi_vi_pham, "so_bb": 0, "tien": 0, "cho_ky": 0})
        if bb.co_hieu_luc:
            o["so_bb"] += 1
            o["tien"] += bb.so_tien or 0
        elif bb.trang_thai == TrangThaiBienBan.CHO_KY:
            o["cho_ky"] += 1
    tong = sorted(tong_theo_nv.values(), key=lambda x: (-x["tien"], x["nv"].ho_ten))
    nhan_vien = []
    if current_user.la_quan_ly:
        ids = {r[0] for r in _q_pham_vi().with_entities(BienBanPhat.nguoi_vi_pham_id).distinct()}
        nhan_vien = NguoiDung.query.filter(NguoiDung.id.in_(ids)).order_by(NguoiDung.ho_ten).all() if ids else []
    tham_so = {k: v for k, v in f.items() if v}
    return render_template(
        "bien_ban_list.html", ds=ds, f=f, tham_so=tham_so, tong=tong, nhan_vien=nhan_vien,
        TrangThaiBienBan=TrangThaiBienBan, duoc_lap=current_user.la_quan_ly,
        tong_tien=sum(x["tien"] for x in tong), thang_nay=ngay_vn_hien_tai().strftime("%Y-%m"),
    )


@bp.route("/moi", methods=["GET", "POST"])
@login_required
def moi():
    if not current_user.la_quan_ly:
        abort(403)
    ds_loi = LoiViPham.query.filter_by(dang_dung=True).order_by(LoiViPham.ten).all()

    if request.method == "GET":
        return render_template("bien_ban_form.html", nhan_vien=_nhan_vien_lap_duoc(), ds_loi=ds_loi,
                               hom_nay=ngay_vn_hien_tai(), chon_nv=request.args.get("nv", type=int))

    def loi(msg):
        flash(msg, "error")
        return redirect(url_for("bien_ban.moi"))

    nv = db.session.get(NguoiDung, request.form.get("nguoi_vi_pham_id", type=int) or 0)
    if not duoc_lap_cho(nv):
        return loi("Chưa chọn nhân viên, hoặc bạn không có quyền lập biên bản cho người này.")
    loi_id = request.form.get("loi_id") or ""
    loi_obj = db.session.get(LoiViPham, int(loi_id)) if loi_id.isdigit() else None
    ten_loi = loi_obj.ten if loi_obj else (request.form.get("ten_loi_khac") or "").strip()
    if not ten_loi:
        return loi("Chưa chọn lỗi vi phạm (hoặc ghi tên lỗi nếu chọn \"Lỗi khác\").")
    noi_dung = (request.form.get("noi_dung") or "").strip()
    if not noi_dung:
        return loi("Chưa ghi diễn biến sự việc.")
    try:
        ngay_vp = date.fromisoformat(request.form.get("ngay_vi_pham") or "")
    except ValueError:
        return loi("Ngày vi phạm không hợp lệ.")
    if ngay_vp > ngay_vn_hien_tai():
        return loi("Ngày vi phạm không được ở tương lai.")
    so_tien = _doc_tien(request.form.get("so_tien"))
    if so_tien is None:
        return loi("Số tiền phạt không hợp lệ — chỉ nhập số, VD 200.000.")

    trung = BienBanPhat.query.filter(
        BienBanPhat.nguoi_lap_id == current_user.id, BienBanPhat.nguoi_vi_pham_id == nv.id,
        BienBanPhat.ten_loi == ten_loi[:200], BienBanPhat.ngay_vi_pham == ngay_vp,
        BienBanPhat.tao_luc >= gio_vn_hien_tai() - timedelta(minutes=3)).first()
    if trung:
        flash("Biên bản này vừa được lập rồi — không lập thêm lần nữa.", "info")
        return redirect(url_for("bien_ban.chi_tiet", id=trung.id))

    chu_ky = _doc_chu_ky()
    if not chu_ky:
        return loi("Cần ký tên người lập biên bản.")

    bb = BienBanPhat(nguoi_vi_pham_id=nv.id, nguoi_lap_id=current_user.id, loi_id=loi_obj.id if loi_obj else None,
                     ten_loi=ten_loi[:200], ngay_vi_pham=ngay_vp, noi_dung=noi_dung, so_tien=so_tien,
                     trang_thai=TrangThaiBienBan.CHO_KY, chu_ky_lap=_luu_chu_ky(chu_ky))
    db.session.add(bb)
    db.session.flush()
    for f in [f for f in request.files.getlist("dinh_kem") if f and f.filename]:
        loai = services.phan_loai(f.filename, f.mimetype) or LoaiDinhKem.FILE
        duong_dan, kich_thuoc = services.luu_file(f, "bien-ban")
        db.session.add(DinhKemBienBan(bien_ban_id=bb.id, loai=loai, duong_dan=duong_dan,
                                      ten_goc=f.filename[:255], kich_thuoc=kich_thuoc, mime=f.mimetype))
    db.session.flush()
    db.session.refresh(bb)
    _sinh_pdf(bb)
    db.session.commit()

    services.bao_bien_ban_moi(bb)
    db.session.commit()
    flash(f"Đã lập biên bản {bb.so_bien_ban} — đã báo Zalo để {nv.ho_ten} vào ký xác nhận.", "success")
    return redirect(url_for("bien_ban.chi_tiet", id=bb.id))


@bp.route("/<int:id>")
@login_required
def chi_tiet(id):
    bb = db.session.get(BienBanPhat, id) or abort(404)
    if not duoc_xem(bb):
        abort(403)
    return render_template(
        "bien_ban_detail.html", bb=bb, TrangThaiBienBan=TrangThaiBienBan,
        duoc_ky=bb.trang_thai == TrangThaiBienBan.CHO_KY and bb.nguoi_vi_pham_id == current_user.id,
        duoc_huy=duoc_huy(bb), duoc_ghi_khong_ky=duoc_ghi_khong_ky(bb),
        duoc_xoa=current_user.la_admin_thuan,
    )


@bp.route("/<int:id>/ky", methods=["POST"])
@login_required
def ky(id):
    bb = db.session.get(BienBanPhat, id) or abort(404)
    if bb.nguoi_vi_pham_id != current_user.id:
        abort(403)
    if bb.trang_thai != TrangThaiBienBan.CHO_KY:
        flash("Biên bản này không còn ở trạng thái chờ ký.", "error")
        return redirect(url_for("bien_ban.chi_tiet", id=id))
    chu_ky = _doc_chu_ky()
    if not chu_ky:
        flash("Cần ký tên trước khi xác nhận.", "error")
        return redirect(url_for("bien_ban.chi_tiet", id=id))
    bb.chu_ky_nv = _luu_chu_ky(chu_ky)
    bb.y_kien_nv = (request.form.get("y_kien_nv") or "").strip() or None
    bb.ky_luc = gio_vn_hien_tai()
    bb.trang_thai = TrangThaiBienBan.DA_KY
    _sinh_pdf(bb)
    db.session.commit()
    services.bao_bien_ban_da_ky(bb)
    db.session.commit()
    flash("Đã ký xác nhận biên bản.", "success")
    return redirect(url_for("bien_ban.chi_tiet", id=id))


@bp.route("/<int:id>/khong-ky", methods=["POST"])
@login_required
def khong_ky(id):
    bb = db.session.get(BienBanPhat, id) or abort(404)
    if not duoc_ghi_khong_ky(bb):
        abort(403)
    bb.trang_thai = TrangThaiBienBan.KHONG_KY
    bb.xu_ly_boi_id = current_user.id
    bb.xu_ly_luc = gio_vn_hien_tai()
    bb.ly_do_xu_ly = (request.form.get("ghi_chu") or "").strip() or None
    _sinh_pdf(bb)
    db.session.commit()
    services.bao_bien_ban_xu_ly(bb)
    db.session.commit()
    flash("Đã ghi nhận nhân viên không ký — biên bản vẫn có hiệu lực.", "success")
    return redirect(url_for("bien_ban.chi_tiet", id=id))


@bp.route("/<int:id>/huy", methods=["POST"])
@login_required
def huy(id):
    bb = db.session.get(BienBanPhat, id) or abort(404)
    if not duoc_huy(bb):
        abort(403)
    ly_do = (request.form.get("ly_do") or "").strip()
    if not ly_do:
        flash("Cần ghi lý do huỷ biên bản.", "error")
        return redirect(url_for("bien_ban.chi_tiet", id=id))
    bb.trang_thai = TrangThaiBienBan.DA_HUY
    bb.xu_ly_boi_id = current_user.id
    bb.xu_ly_luc = gio_vn_hien_tai()
    bb.ly_do_xu_ly = ly_do
    _sinh_pdf(bb)
    db.session.commit()
    services.bao_bien_ban_xu_ly(bb)
    db.session.commit()
    flash(f"Đã huỷ biên bản {bb.so_bien_ban}.", "success")
    return redirect(url_for("bien_ban.chi_tiet", id=id))


@bp.route("/<int:id>/xoa", methods=["POST"])
@login_required
def xoa(id):
    if not current_user.la_admin_thuan:
        abort(403)
    bb = db.session.get(BienBanPhat, id) or abort(404)
    so = bb.so_bien_ban
    for p in [bb.duong_dan_pdf, bb.chu_ky_lap, bb.chu_ky_nv] + [d.duong_dan for d in bb.dinh_kem]:
        services._xoa_file_tuong_doi(p)
    db.session.delete(bb)
    db.session.commit()
    flash(f"Đã xoá vĩnh viễn biên bản {so}.", "success")
    return redirect(url_for("bien_ban.danh_sach"))


@bp.route("/xuat-excel")
@login_required
def xuat_excel():
    if not current_user.la_quan_ly:
        abort(403)
    f, q = _bo_loc()
    wb = Workbook()
    ws = wb.active
    ws.title = "Biên bản phạt"
    ws.append(["Số BB", "Ngày vi phạm", "Mã NV", "Nhân viên", "Bộ phận", "Lỗi vi phạm", "Diễn biến",
               "Tiền phạt", "Trạng thái", "Người lập", "Ngày lập"])
    for o in ws[1]:
        o.font = Font(bold=True, color="FFFFFF")
        o.fill = PatternFill("solid", fgColor="1F4E79")
        o.alignment = Alignment(horizontal="center", vertical="center")
    for bb in q.all():
        nv = bb.nguoi_vi_pham
        ws.append([bb.so_bien_ban, bb.ngay_vi_pham.strftime("%d/%m/%Y"), nv.ma_dinh_danh, nv.ho_ten,
                   nv.ten_cac_bo_phan, bb.ten_loi, bb.noi_dung, bb.so_tien or 0, bb.ten_trang_thai,
                   bb.nguoi_lap.ho_ten, bb.tao_luc.strftime("%d/%m/%Y %H:%M") if bb.tao_luc else ""])
        ws.cell(ws.max_row, 8).number_format = "#,##0"
    for cot, rong in zip("ABCDEFGHIJK", (13, 13, 10, 24, 18, 28, 40, 12, 13, 22, 16)):
        ws.column_dimensions[cot].width = rong
    ws.freeze_panes = "A2"
    tep = BytesIO()
    wb.save(tep)
    tep.seek(0)
    ten = f"bien-ban-phat-{f['thang'] or ngay_vn_hien_tai().strftime('%Y-%m-%d')}.xlsx"
    return send_file(tep, as_attachment=True, download_name=ten,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ------------------------------------------------------------ danh mục lỗi
@bp.route("/danh-muc-loi", methods=["GET", "POST"])
@login_required
def danh_muc_loi():
    if not current_user.la_admin_sep:
        abort(403)
    if request.method == "POST":
        lid = request.form.get("id", type=int)
        l = db.session.get(LoiViPham, lid) if lid else LoiViPham()
        if lid and not l:
            abort(404)
        ten = " ".join((request.form.get("ten") or "").split())
        muc = _doc_tien(request.form.get("muc_phat"))
        if not ten or muc is None:
            flash("Cần nhập tên lỗi và mức phạt hợp lệ (chỉ số, có thể 0).", "error")
            return redirect(url_for("bien_ban.danh_muc_loi"))
        trung = LoiViPham.query.filter(db.func.lower(LoiViPham.ten) == ten.lower()).first()
        if trung and trung.id != l.id:
            flash(f"Lỗi \"{ten}\" đã có trong danh mục.", "error")
            return redirect(url_for("bien_ban.danh_muc_loi"))
        l.ten = ten[:200]
        l.muc_phat = muc
        l.mo_ta = (request.form.get("mo_ta") or "").strip() or None
        l.dang_dung = request.form.get("dang_dung") == "on" if lid else True
        if not lid:
            db.session.add(l)
        db.session.commit()
        flash("Đã lưu lỗi vi phạm.", "success")
        return redirect(url_for("bien_ban.danh_muc_loi"))
    ds = LoiViPham.query.order_by(LoiViPham.dang_dung.desc(), LoiViPham.ten).all()
    so_lan = dict(db.session.query(BienBanPhat.loi_id, db.func.count(BienBanPhat.id))
                  .filter(BienBanPhat.loi_id.isnot(None)).group_by(BienBanPhat.loi_id).all())
    return render_template("bien_ban_loi.html", ds=ds, so_lan=so_lan)


@bp.route("/danh-muc-loi/<int:id>/xoa", methods=["POST"])
@login_required
def xoa_loi(id):
    if not current_user.la_admin_sep:
        abort(403)
    l = db.session.get(LoiViPham, id) or abort(404)
    BienBanPhat.query.filter_by(loi_id=l.id).update({"loi_id": None})
    db.session.delete(l)
    db.session.commit()
    flash(f"Đã xoá lỗi \"{l.ten}\" khỏi danh mục (biên bản cũ vẫn giữ nguyên tên lỗi).", "success")
    return redirect(url_for("bien_ban.danh_muc_loi"))
